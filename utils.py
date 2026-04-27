import os
import re
import random
import time
from groq import Groq
import PyPDF2
import docx
import io
import json
from fpdf import FPDF
import matplotlib.pyplot as plt
import tempfile

client = None


def set_groq_api_key(api_key):
    """Configure the Groq client from a user-provided API key."""
    global client
    cleaned_key = (api_key or "").strip()
    if not cleaned_key:
        client = None
        return None

    client = Groq(api_key=cleaned_key)
    return client


def get_groq_client():
    """Return the configured Groq client or raise a helpful error."""
    if client is None:
        raise RuntimeError("Groq API key is not configured. Enter it in the sidebar first.")
    return client

def extract_text_from_file(uploaded_file):
    """Extracts text from uploaded file based on its type."""
    file_type = uploaded_file.name.split('.')[-1].lower()
    
    if file_type == 'txt':
        return uploaded_file.read().decode("utf-8")
    
    elif file_type == 'pdf':
        pdf_reader = PyPDF2.PdfReader(uploaded_file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text
    
    elif file_type == 'docx':
        doc = docx.Document(uploaded_file)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text
    
    return ""

BLOOMS_TAXONOMY = {
    "L1": {
        "level": "Remember",
        "verbs": ["Define", "List", "Identify", "Recall", "Name", "State", "Label", "Write", "Recognize", "Match", "Memorize", "What", "Which", "Who", "When", "Where"]
    },
    "L2": {
        "level": "Understand",
        "verbs": ["Explain", "Describe", "Summarize", "Interpret", "Classify", "Compare", "Discuss", "Illustrate", "Differentiate", "Paraphrase", "How", "Why", "Elaborate", "Contrast", "Outline"]
    },
    "L3": {
        "level": "Apply",
        "verbs": ["Calculate", "Solve", "Demonstrate", "Use", "Implement", "Compute", "Simulate", "Apply", "Execute", "Operate", "Show", "Solve", "Practise"]
    },
    "L4": {
        "level": "Analyze",
        "verbs": ["Analyze", "Examine", "Test", "Investigate", "Categorize", "Infer", "Distinguish", "Diagnose", "Breakdown", "Correlation", "Diagram"]
    },
    "L5": {
        "level": "Evaluate",
        "verbs": ["Evaluate", "Justify", "Critique", "Validate", "Assess", "Recommend", "Defend", "Optimize", "Prioritize", "Verify", "Judge", "Appraise"]
    },
    "L6": {
        "level": "Create",
        "verbs": ["Design", "Develop", "Construct", "Formulate", "Propose", "Create", "Build", "Model", "Invent", "Plan", "Integrate", "Compose", "Generate", "Modify"]
    }
}

def identify_blooms_level(question_text):
    """
    Identifies the Bloom's Level (L1-L6) based on the first word or keywords in the question.
    """
    if not question_text:
        return "L1"
        
    # Standardize text for mapping
    clean_q = question_text.strip().lstrip('Q0123456789. ').strip()
    first_word = clean_q.split()[0].rstrip(',:?').capitalize() if clean_q.split() else ""
    
    # Check verb matches
    for level, data in BLOOMS_TAXONOMY.items():
        if first_word in data['verbs']:
            return level
            
    # Substring check if first word didn't match perfectly
    for level in ["L6", "L5", "L4", "L3", "L2", "L1"]: # Check higher levels first
        for verb in BLOOMS_TAXONOMY[level]['verbs']:
            if verb.lower() in clean_q.lower()[:20]: # Check beginning of question
                return level
                
    return "L1" # Default fallback

def get_starter_word(text):
    """Extracts the first word of a question for diversity checking."""
    if not text: return ""
    clean_q = text.strip().lstrip('Q0123456789. ').strip()
    return clean_q.split()[0].rstrip(',:?').capitalize() if clean_q.split() else ""

def fix_diversity(questions, used_starters):
    """Ensures every question in the set starts with a unique word by rephrasing if necessary."""
    for q in questions:
        current_starter = get_starter_word(q.get('question', ''))
        if not current_starter: continue
        
        if current_starter in used_starters:
            # Rephrase using another verb from same Blooms level
            blooms = q.get('blooms', 'L1')
            verbs = BLOOMS_TAXONOMY.get(blooms, {}).get("verbs", [])
            # Add some universal fallbacks
            universal = ["How", "Why", "What", "Which", "Discuss", "State", "Describe"]
            combined_verbs = list(dict.fromkeys(verbs + universal))
            random.shuffle(combined_verbs)
            
            for v in combined_verbs:
                v_cap = v.capitalize()
                if v_cap not in used_starters:
                    orig = q.get('question', '')
                    clean_q = orig.strip().lstrip('Q0123456789. ').strip()
                    parts = clean_q.split()
                    if parts:
                        parts[0] = v_cap
                        q['question'] = " ".join(parts)
                        current_starter = v_cap
                        break
        used_starters.add(current_starter)
    return questions

def extract_topic_keywords(text):
    """Simple extraction of main nouns/topics from a question to avoid repetition."""
    if not text: return []
    # Remove standard question starters and common stop words
    ignore = ["explain", "describe", "identify", "what", "which", "how", "why", "discuss", "compare", "contrast", "significance", "context", "real-world", "applications", "context", "importance", "role", "principles", "concept", "context-free"]
    clean = re.sub(r'[^a-zA-Z0-9\s]', '', text.lower())
    words = clean.split()
    keywords = [w for w in words if len(w) > 3 and w not in ignore]
    return keywords

def sanitize_text(text):
    """Clean text for FPDF compatibility (Standard fonts only support Latin-1)."""
    if not text:
        return ""
    # Map common unicode characters to latin-1 or similar
    chars = {
        '\u2013': '-', '\u2014': '-', 
        '\u2018': "'", '\u2019': "'", 
        '\u201c': '"', '\u201d': '"',
        '\u2022': '*', '\u2026': '...'
    }
    for u_char, r_char in chars.items():
        text = text.replace(u_char, r_char)
    return text.encode('latin-1', 'replace').decode('latin-1')

def extract_units(syllabus_text):
    """
    Universal syllabus extractor for Unit 1, Unit 2 and Unit 3.
    Handlers various separators like 'UNIT-I', 'UNIT - II', 'UNIT – III', etc.
    """
    # More flexible patterns that handle hyphens, dashes and extra spaces
    # Matches 'UNIT' or 'MODULE' or 'PART' followed by optional non-word chars, then the Roman/Arabic number
    # Negative lookahead (?![I\d]) prevents matching UNIT I as the start of UNIT II, but allows following words like SOFTWARE.
    u1_patterns = [r"(?:UNIT|MODULE|PART)[\s\W]*(?:I|1)(?![I\d])", r"(?:^|[\n\r])\s*(?:UNIT|MODULE|PART)\s*1\s*[:\-]"]
    u2_patterns = [r"(?:UNIT|MODULE|PART)[\s\W]*(?:II|2)(?![I\d])", r"(?:^|[\n\r])\s*(?:UNIT|MODULE|PART)\s*2\s*[:\-]"]
    u3_patterns = [r"(?:UNIT|MODULE|PART)[\s\W]*(?:III|3)(?![I\d])", r"(?:^|[\n\r])\s*(?:UNIT|MODULE|PART)\s*3\s*[:\-]"]
    
    # Improved terminators to match Unit 4, 5 and other common end-of-unit markers
    terminators = r"(?:UNIT|MODULE|PART)[\s\W]*(?:IV|4|V|5|VI|6)|Total\s*Hours|TEXT\s*BOOKS|REFERENCES|Note:|Total\s*:|Outcome|Syllabus|Objective"

    unit1 = ""
    unit2 = ""
    unit3 = ""
    
    all_positions = []
    
    # Use re.IGNORECASE and re.MULTILINE for all matches
    for pattern in u1_patterns:
        matches = list(re.finditer(pattern, syllabus_text, re.IGNORECASE | re.MULTILINE))
        for match in matches: all_positions.append(('u1', match.start(), match.end()))
    
    for pattern in u2_patterns:
        matches = list(re.finditer(pattern, syllabus_text, re.IGNORECASE | re.MULTILINE))
        for match in matches: all_positions.append(('u2', match.start(), match.end()))
    
    for pattern in u3_patterns:
        matches = list(re.finditer(pattern, syllabus_text, re.IGNORECASE | re.MULTILINE))
        for match in matches: all_positions.append(('u3', match.start(), match.end()))
    
    all_positions.sort(key=lambda x: x[1])
    
    for i, (unit_type, start_pos, _) in enumerate(all_positions):
        end_pos = len(syllabus_text)
        if i + 1 < len(all_positions):
            end_pos = all_positions[i+1][1]
        
        block = syllabus_text[start_pos:end_pos]
        term_match = re.search(terminators, block, re.IGNORECASE)
        # Only terminate if the match is NOT at the very start (preventing self-termination)
        if term_match and term_match.start() > 5:
            block = block[:term_match.start()]
            
        if unit_type == 'u1' and not unit1:
            unit1 = block.strip()
        elif unit_type == 'u2' and not unit2:
            unit2 = block.strip()
        elif unit_type == 'u3' and not unit3:
            unit3 = block.strip()
            
    return unit1, unit2, unit3

def extract_cos(syllabus_text):
    """Extracts Course Outcomes (COs) from the syllabus text, including Bloom's levels if present."""
    cos = []
    
    # Updated pattern to capture Bloom's level if it exists (e.g., [L2])
    # Matches: CO1: Description... [L2]
    strict_pattern = r"(?:^|[\n\r])\s*(CO\d+)\s*[:\-\.\s]*\s*(.+?)(?=\s*\[L\d\]|Mapping|\n\s*Mapping|PO\d|[\n\r]|$)(?:\s*\[(L\d)\])?"
    matches = list(re.finditer(strict_pattern, syllabus_text, re.IGNORECASE | re.MULTILINE))
    
    if matches:
        for match in matches:
            co_id = match.group(1).upper().replace(" ", "")
            co_desc = match.group(2).strip()
            blooms_level = match.group(3) if match.group(3) else None
            
            # Clean up potential trailing brackets or mapping text
            co_desc = re.sub(r'\[.*?\]', '', co_desc).strip()
            if "Mapping of" in co_desc:
                co_desc = co_desc.split("Mapping of")[0].strip()
            
            if co_desc and len(co_desc) > 5:
                cos.append({"id": co_id, "description": co_desc, "blooms": blooms_level})
    
    # Fallback to broader pattern
    if not cos:
        # Matches CO1 ... until next CO, UNIT, or common header
        co_pattern = r"(CO\d+|Course\s*Outcome\s*\d+)[\s:-]+(.*?)(?=CO\d+|Course\s*Outcome\s*\d+|UNIT|Total|REFERENCE|Mapping|PO\d|PSO\d|$)"
        matches = re.finditer(co_pattern, syllabus_text, re.IGNORECASE | re.DOTALL)
        for match in matches:
            co_id = match.group(1).upper().replace(" ", "")
            co_desc = match.group(2).strip()
            
            # Try to find bloom's level: [L2], (L2), (Level 2)
            bloom_match = re.search(r'(?:\[|\()(?:L|Level\s*)(\d)(?:\]|\))', co_desc, re.IGNORECASE)
            blooms_level = f"L{bloom_match.group(1)}" if bloom_match else None
            
            # Clean up the description
            co_desc = re.sub(r'(?:\[|\()(?:L|Level\s*)\d(?:\]|\))', '', co_desc, flags=re.IGNORECASE).strip()
            
            if "Mapping of" in co_desc:
                co_desc = co_desc.split("Mapping of")[0].strip()
            if co_desc and len(co_desc) > 10:
                cos.append({"id": co_id, "description": co_desc, "blooms": blooms_level})

    # Fallback 2: Look for "Course Outcomes" heading
    if not cos:
        co_section = re.search(r"Course\s*Outcomes.*?(?=\n\s*\n|UNIT|REFERENCE|Mapping|PO\d|PSO\d|$)", syllabus_text, re.IGNORECASE | re.DOTALL)
        if co_section:
            # Match items starting with CO1, 1., *, etc.
            items = re.findall(r"(?:^|[\n\r])\s*(?:\d+\.|\*|-|CO\d+)\s*(.*)", co_section.group(0), re.MULTILINE)
            for i, item in enumerate(items):
                clean_item = item.strip()
                # Check for Bloom's level patterns: [L2], (L2), (Level 2), Bloom's Level 2
                bloom_match = re.search(r'(?:\[|\()(?:L|Level\s*)(\d)(?:\]|\))', clean_item, re.IGNORECASE)
                blooms_level = f"L{bloom_match.group(1)}" if bloom_match else None
                
                # Clean up the description
                clean_item = re.sub(r'(?:\[|\()(?:L|Level\s*)\d(?:\]|\))', '', clean_item, flags=re.IGNORECASE).strip()
                
                if "Mapping" in clean_item:
                    clean_item = clean_item.split("Mapping")[0].strip()
                if clean_item and len(clean_item) > 10:
                    cos.append({"id": f"CO{i+1}", "description": clean_item, "blooms": blooms_level})
    
    if not cos:
        return [{"id": "CO1", "description": "Successfully complete all units.", "blooms": "L2"}]
    
    # Remove duplicates preserving order
    unique_cos = []
    seen_ids = set()
    for c in cos:
        if c['id'] not in seen_ids and len(c['description']) > 10:
            seen_ids.add(c['id'])
            unique_cos.append(c)
            
    return unique_cos

def split_syllabus_by_topics(text, n=3):
    """
    Splits the syllabus text into n roughly equal parts based on comma or hyphenated topics.
    Useful for ensuring question paper coverage.
    """
    if not text:
        return [""] * n
        
    # More aggressive removal of unit header
    # 1. Remove "UNIT" followed by Roman/Arabic index (e.g., UNIT II, UNIT-2)
    content = re.sub(r"^UNIT[\s\W]*(?:[IV\d]+)", "", text, flags=re.IGNORECASE).strip()
    
    # 2. Remove the unit title and hours (e.g., WORD LEVEL ANALYSIS 9 Hrs)
    # Use DOTALL to handle cases where the title might span multiple lines
    content = re.sub(r"^(?:.*?)(?:\d+\s*(?:Hrs|Hours))", "", content, flags=re.IGNORECASE | re.DOTALL).strip()
    
    # Split by common delimiters: comma, semicolon, em-dash
    items = [item.strip() for item in re.split(r",|;|–", content) if item.strip()]
    
    # Fallback: If we only got 1 item, try splitting by hyphens
    if len(items) <= 1:
        # Only split by hyphens if there are several of them (likely separators)
        if content.count('-') >= 2:
            items = [item.strip() for item in re.split(r"-", content) if item.strip()]
        
    if not items or (len(items) == 1 and not items[0]):
        # Final fallback
        items = [content] if content else ["Syllabus Content"]

    # If still fewer than n, pad with placeholders
    while len(items) < n:
        items.append("...")
    
    # Divide items into n groups
    k, m = divmod(len(items), n)
    parts = []
    start = 0
    for i in range(n):
        end = start + k + (1 if i < m else 0)
        part_items = items[start:end]
        if part_items:
            # Re-join using commas for display
            parts.append(", ".join(part_items))
        else:
            parts.append("")
        start = end
    
    return parts

def generate_questions(u1, u2, u3, exam_type, cos, difficulty_config, blooms_taxonomy, subject_name="Subject", exclude_questions=None):
    """Generates questions using Groq LLM with aggressive prompt optimization and CO mapping."""
    client = get_groq_client()
    sys_random = random.SystemRandom()
    run_nonce = f"{int(time.time() * 1000)}-{sys_random.randint(100000, 999999)}"
    
    # 1. Prepare CO context for prompt
    co_context = "\n".join([f"- {c['id']}: {c['description']} (Target: {c.get('blooms', 'Any')})" for c in cos])

    # 1. Compact Difficulty Definitions
    DIFFICULTY_RULES = f"""
    🟢 EASY: 60% L1 (Remember), 40% L2 (Understand).
    🟡 MEDIUM: 20% L1, 30% L2, 40% L3 (Apply), 10% L4 (Analyze).
    🔴 HARD: 10% L2, 20% L3, 40% L4, 20% L5 (Evaluate), 10% L6 (Create).
    
    CAT1: Unit 1 Difficulty: {difficulty_config.get('u1', 'Medium')}
    CAT2: Unit 2 Difficulty: {difficulty_config.get('u2', 'Medium')}, Unit 3 Difficulty: {difficulty_config.get('u3', 'Medium')}
    """

    # 2. Compact Rules
    BLOOMS_KEYWORDS = "\n".join([f"- {level} ({data['level']}): {', '.join(data['verbs'])}" for level, data in blooms_taxonomy.items()])

    # Avoidance logic
    avoidance_instr = ""
    if exclude_questions and len(exclude_questions) > 0:
        # Take only the most recent ones to avoid token bloat, e.g., last 30 questions
        recent_excludes = exclude_questions[-30:]
        avoidance_instr = f"""
        CRITICAL: DO NOT REGENERATE OR REPEAT ANY OF THE FOLLOWING QUESTIONS:
        {chr(10).join([f"- {q}" for q in recent_excludes])}
        
        Ensure the new questions cover DIFFERENT topics or perspectives within the same syllabus units.
        """

    BASE_RULES = f"""
    ROLE: Expert Professor. Generate valid JSON only.
    CONSTRAINT: Use ONLY provided Syllabus text. NO outside topics.
    CO MAPPING: Use this list:
    {co_context}
    
    BLOOMS TAXONOMY MAPPING (STRICT):
    Use these specific keywords/verbs to start questions for each level:
    {BLOOMS_KEYWORDS}

    {avoidance_instr}

    DIVERSITY & PHRASING (STRICT):
    - **UNIQUE STARTERS**: Every question MUST start with a different word. No two questions in the entire set can share the same opening word.
    - **VARY THE STYLE**: Use the Bloom's verbs provided above.
    - **STRUCTURES**: 
        - "What are the significant impacts of [X] on [Y]?"
        - "Which factor plays a critical role in [Z]?"
        - "Identify the primary methodology used for..."
        - "Compare [X] and [Y] in the context of..."
        - "How does [X] handle the challenge of [Y]?"
        - "Describe the process of..."
    - **NO COPYING**: Do NOT just copy the examples above. Use them as inspiration for variety.
    - **ZERO TOLERANCE**: If any 2 questions start with the same word, it is a failure.
    
    SELF-CONTAINED (STRICT):
    - **NO "GIVEN X"**: Do NOT use words like "given algorithm", "following table", or "provided code" unless YOU output that code/table in the question.
    - Questions must be complete and understandable on their own.
    """

    # 3. Dynamic Syllabus Context Construction
    def clean_text(t): return t.strip()[:15000]
    context_a = ""
    context_b = ""
    
    if exam_type == "CAT1":
        u1_text = f"Unit 1 Content:\n{clean_text(u1)}"
        context_a = u1_text
        context_b = u1_text
    else:
        u2_text = f"Unit 2 Content:\n{clean_text(u2)}"
        u3_text = f"Unit 3 Content:\n{clean_text(u3)}"
        context_a = f"{u2_text}\n---\n{u3_text}"
        context_b = f"{u2_text}\n---\n{u3_text}"

    # 4. MCQ Specific Rules
    MCQ_RULES = f"""
    TASK: Generate Part A (10 MCQs).
    STRUCTURE: "part_a": [{{ "number": 1, "unit": "...", "co": "CO1", "blooms": "L1", "question": "...", "options": {{ "A": "...", "B": "...", "C": "...", "D": "..."}}, "answer": "..." }}]
    
    DISTRIBUTION:
    - CAT1: Q1-10 from Unit 1.
    - CAT2: Q1-5 from Unit 2, Q6-10 from Unit 3.
    - BLOOMS MIX: Strictly follow the {difficulty_config.get('u1', 'Medium')} distribution. Ensure at least 3-4 questions are L2/L3.
    
    PHRASING DIVERSITY:
    - **NO REPETITION**: Every question MUST start with a unique word.
    - **VARIED INQUIRY**: Use "What...", "Which...", "How...", "Identify...", "Define...", "When...", "Why...".
    - **UNPREDICTABLE**: Do not follow a pattern. If Q1 starts with "Which", Q2 must start with something else like "How".
    
    OPTIONS DIVERSITY (ZERO TOLERANCE):
    - **NO SHARED STARTING PHRASES**: Every option (A, B, C, D) MUST start with a unique word. 
    - **AVOID SUBJECT REPETITION**: Do not start all options with the same subject (e.g., "The system...", "The system...").
    - **NO REPETITIVE PHRASES**: Do not start options with "A...", "The...", "By...", "To...", "In..." repeatedly.
    - **VARY GRAMMAR**: If Option A starts with a noun, Option B must start with a verb, Option C with a gerund, etc.
    - **CONTENT VARIETY**: Ensure each option represents a distinct concept or approach.
    - **REWRITE IF REPETITIVE**: If you find yourself repeating the first 2-3 words in every option, move that phrase into the question stem and keep options concise and distinct.
    - **REJECTION CRITERIA**: Any question where 2 or more options start with the same word is considered a FAILURE.
    """

    # 5. Subjective Specific Rules
    SUBJECTIVE_RULES = f"""
    TASK: Generate Part B (Subjective) with Answers.
    STRUCTURE: "part_b": [{{ "number": 11, "question": "...", "unit": "...", "co": "...", "blooms": "...", "answer": "..." }}]
    Marks: 10 Marks each.
    
    ANSWER GENERATION REQUIREMENTS:
    - For EVERY Part B question, generate exactly 7 to 8 distinct technical points.
    - DO NOT use sub-headings or bold titles (e.g., NO "Definition:", NO "**Purpose**:").
    - Start each point directly with a technical explanation.
    - Each point MUST be a complete, accurate sentence.
    - Format: 1. [Technical sentence] \n 2. [Technical sentence]...
    
    PHRASING CONSTRAINTS:
    - **UNIQUE OPENINGS**: Every question MUST start with a unique word.
    - **DIVERSE STYLE**: Use a natural variety: "What...", "Which...", "How...", "Explain...", "Describe...", "Compare...", "Contrast...".
    - **NO TEMPLATES**: Avoid using the same sentence pattern across questions.
    - **BLOOMS LEVEL**: Part B should target higher levels (L3, L4, L5, L6).
    """

    paper = {"exam_type": exam_type, "part_a": [], "part_b": []}
    used_starters = set()
    
    # Track starters and keywords from excluded questions
    if exclude_questions:
        for old_q in exclude_questions:
            word = get_starter_word(old_q)
            if word: used_starters.add(word)
    
    # 2. Extract keywords from previously generated questions to avoid topic overlap
    forbidden_topics = []
    if exclude_questions:
        for old_q in exclude_questions:
            forbidden_topics.extend(extract_topic_keywords(old_q))
    forbidden_topics = list(set(forbidden_topics))[:40] # Capped for prompt length
    
    forbidden_topics_str = ", ".join(forbidden_topics)
    
    # helper for attribute rotation
    def rotate_attributes(index, available_cos, difficulty):
        # CO Rotation
        if not available_cos: co = "CO1"
        else: co = available_cos[index % len(available_cos)]['id']
        
        # Blooms Rotation based on Difficulty
        if difficulty == "Easy":
            blooms_opts = ["L1", "L2"]
        elif difficulty == "Medium":
            blooms_opts = ["L2", "L3"]
        else: # Hard
            blooms_opts = ["L3", "L4", "L5"]
            
        blooms = blooms_opts[index % len(blooms_opts)]
        return co, blooms
    
    # 6. Execution with Delays
    
    # Final Diversity Footer
    DIVERSITY_FOOTER = """
    CRITICAL CONSTRAINTS:
    - NO REPETITION: Every question MUST start with a unique word.
    - STARTER VARIETY: Actively use "What", "Which", "How", "Identify", "Compare", "Explain".
    - NATURAL LANGUAGE: Do not use the same template for every question.
    - TOPIC DIVERSITY: Within a syllabus segment, explore different concepts, not just the first one.
    - EVERY OPENING MUST BE DISTINCT.
    """

    # --- PART A GENERATION ---
    variation_seed = sys_random.randint(1, 1000000)
    prompt_a = f"{BASE_RULES}\n{DIFFICULTY_RULES}\n{MCQ_RULES}\nSYLLABUS:\n{context_a}\n{DIVERSITY_FOOTER}\nVARIATION_SEED: {variation_seed}\nRUN_NONCE: {run_nonce}\nJSON:"
    try:
        completion_a = client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[{"role": "user", "content": prompt_a}],
            response_format={"type": "json_object"},
            temperature=0.85,
            top_p=0.95,
            max_tokens=2048
        )
        data_a = json.loads(completion_a.choices[0].message.content)
        if 'part_a' in data_a: 
            paper['part_a'] = data_a['part_a'][:10]
            # Track starters
            for q in paper['part_a']:
                word = get_starter_word(q.get('question'))
                if word: used_starters.add(word)
    except Exception as e:
        paper['error'] = str(e)

    # --- PART B GENERATION (Buffered/Split for CAT2) ---
    def get_part_b(chunk_desc, count, syllabus_context, current_used, forbidden_topics_list):
        seed = sys_random.randint(1, 1000000)
        forbidden_s = ", ".join(list(current_used))
        forbidden_t = ", ".join(forbidden_topics_list)
        
        p_b = f"""
        {BASE_RULES}
        {DIFFICULTY_RULES}
        {SUBJECTIVE_RULES}
        
        SCOPE: {chunk_desc}
        SYLLABUS SEGMENT:
        {syllabus_context}
        
        DIVERSITY CONSTRAINTS:
        - FORBIDDEN_STARTERS: {forbidden_s}
        - FORBIDDEN_TOPICS/KEYWORDS: {forbidden_t}
        - RULE: If the SYLLABUS SEGMENT contains multiple topics, you MUST choose a topic that is NOT in the FORBIDDEN_TOPICS list.
        - RULE: Do NOT center the question around the same main subject as previous versions.
        
        VARIATION_SEED: {seed}
        RUN_NONCE: {run_nonce}
        JSON:"""
        
        res = client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[{"role": "user", "content": p_b}],
            response_format={"type": "json_object"},
            temperature=0.85,
            top_p=0.95,
            max_tokens=3072 # Generous tokens for answers
        )
        content = json.loads(res.choices[0].message.content)
        raw_part_b = content.get('part_b')
        if not isinstance(raw_part_b, list):
            raw_part_b = []
            
        # Filter out None values and ensure they are dicts
        valid_part_b = [q for q in raw_part_b if q and isinstance(q, dict)]
        return valid_part_b[:count]

    # --- PART B GENERATION (Optimized) ---
    import concurrent.futures

    def get_unit_questions(unit_num, syllabus_text, start_q, count, current_used, forbidden_t):
        segments = split_syllabus_by_topics(syllabus_text, n=count)
        
        # Retry mechanism for robust generation
        for attempt in range(2):
            try:
                scope_desc = f"Generate EXACTLY {count} questions (Q{start_q} to Q{start_q + count - 1}) from Unit {unit_num}.\n"
                for i, seg in enumerate(segments):
                    scope_desc += f"- Q{start_q + i} MUST be from this segment: {seg}\n"
                
                qs = get_part_b(scope_desc, count, syllabus_text, current_used, forbidden_t)
                
                # Immediate normalization of basic fields
                for i, q in enumerate(qs):
                    q['number'] = start_q + i
                    q['unit'] = str(unit_num)
                    if 'marks' not in q: q['marks'] = 10 # Default
                
                if len(qs) >= count:
                    return qs
            except Exception as e:
                if attempt == 1: print(f"Part B retry failed: {e}")
                time.sleep(1)
        
        return []

    try:
        if exam_type == "CAT1":
            paper['part_b'] = get_unit_questions(1, u1, 11, 2, used_starters, forbidden_topics)
            fix_diversity(paper['part_b'], used_starters)
        else:
            with concurrent.futures.ThreadPoolExecutor() as executor:
                # Run Unit 2 and Unit 3 in parallel
                future_u2 = executor.submit(get_unit_questions, 2, u2, 11, 3, used_starters.copy(), forbidden_topics.copy())
                # Pass a slightly modified seed/context to avoid same-time collisions if needed
                future_u3 = executor.submit(get_unit_questions, 3, u3, 14, 3, used_starters.copy(), forbidden_topics.copy())
                
                chunk1 = future_u2.result()
                chunk2 = future_u3.result()
            
            # Post-process for diversity after merging
            fix_diversity(chunk1, used_starters)
            fix_diversity(chunk2, used_starters)
            paper['part_b'] = chunk1 + chunk2
            
    except Exception as e:
        if 'error' not in paper: paper['error'] = f"Part B Err: {e}"

    # 7. Post-Processing & Validation
    available_co_ids = [c['id'] for c in cos]

    # Process Part A
    placeholder_warnings = []
    for i, q in enumerate(paper.get('part_a', [])):
        # 0. Silent Fix for "Given/Following" without content
        q_text = q.get('question', '')
        if any(word in q_text.lower() for word in ['given ', 'following ']):
            if not any(marker in q_text for marker in ['```', ':', '{', '[']):
                q['question'] = re.sub(r'\bgiven\s+', '', q['question'], flags=re.IGNORECASE)
                q['question'] = re.sub(r'\bfollowing\s+', '', q['question'], flags=re.IGNORECASE)
                q['question'] = q['question'].capitalize()

        # 1. Enforce Correct Unit/Number
        q['number'] = i + 1
        if exam_type == "CAT1":
            q['unit'] = "1"
            # user might want Q1-10 from Unit 1, but sometimes they want a mix. 
            # For CAT1 we stick to Unit 1 as per conventional rules unless specified.
        else:
            # Always enforce correct unit for CAT2: Q1-5 from Unit 2, Q6-10 from Unit 3
            q['unit'] = "2" if i < 5 else "3"
        
        # 2. Identify Actual Bloom's Level from keywords
        q['blooms'] = identify_blooms_level(q.get('question', ''))
        
        # 3. Validate CO
        if q.get('co') not in available_co_ids:
            q['co'] = available_co_ids[i % len(available_co_ids)]
        
        # 4. ENFORCE MARKS (User Requirement)
        q['marks'] = 1
        # Assuming the LLM provides 'answer' key (e.g., 'A') or we infer it. 
        # Since we want to force distribution, we will shuffle the options and valid answer.
        # However, the LLM usually gives "answer": "A" or similar.
        # We will fully reshuffle options to be safe.
        if 'options' in q and isinstance(q['options'], dict):
            # Extract current valid answer text
            current_ans_key = q.get('answer', 'A').strip().upper()
            # If answer is like "Option A", clean it
            if len(current_ans_key) > 1 and current_ans_key.startswith('OPTION'): 
                current_ans_key = current_ans_key.split()[-1]
            if current_ans_key not in ['A', 'B', 'C', 'D']: current_ans_key = 'A'
            
            correct_text = q['options'].get(current_ans_key, "")
            
            # Get all option texts
            all_opts = [q['options'][k] for k in ['A', 'B', 'C', 'D']]
            random.shuffle(all_opts)
            
            # Reassign to A, B, C, D
            new_options = {k: v for k, v in zip(['A', 'B', 'C', 'D'], all_opts)}
            q['options'] = new_options
            
            # Find new answer key
            new_ans_key = 'A'
            for k, v in new_options.items():
                if v == correct_text:
                    new_ans_key = k
                    break
            q['answer'] = new_ans_key

        # 4. Silent Auto-Cleanup of shared option starters (Longest Common Prefix)
        if 'options' in q and isinstance(q['options'], dict):
            opts_list = [str(q['options'].get(k, "")).strip() for k in ['A', 'B', 'C', 'D']]
            if all(opts_list) and len(opts_list) == 4:
                # Find longest common prefix of words
                word_lists = [opt.split() for opt in opts_list]
                min_len = min(len(wl) for wl in word_lists)
                prefix_len = 0
                for i in range(min_len):
                    first_word = word_lists[0][i].lower()
                    if all(wl[i].lower() == first_word for wl in word_lists):
                        prefix_len += 1
                    else:
                        break
                
                if prefix_len > 0:
                    for k in ['A', 'B', 'C', 'D']:
                        new_val = " ".join(q['options'][k].split()[prefix_len:]).strip()
                        if new_val:
                            q['options'][k] = new_val[0].upper() + new_val[1:] if len(new_val) > 1 else new_val.upper()

    # Process Part B
    for i, q in enumerate(paper.get('part_b', [])):
        # 0. Silent Fix for "Given/Following" without content
        q_text = q.get('question', '')
        if any(word in q_text.lower() for word in ['given ', 'following ']):
            if not any(marker in q_text for marker in ['```', ':', '{', '[']):
                q['question'] = re.sub(r'\bgiven\s+', '', q['question'], flags=re.IGNORECASE)
                q['question'] = re.sub(r'\bfollowing\s+', '', q['question'], flags=re.IGNORECASE)
                q['question'] = q['question'].capitalize()

        # 1. Enforce Correct Unit/Number
        q['number'] = 11 + i
        if exam_type == "CAT1":
            q['unit'] = "1"
        else:
            # Always enforce correct unit for CAT2: Q11-13 from Unit 2, Q14-16 from Unit 3
            q['unit'] = "2" if i < 3 else "3"
        
        # 2. Identify Actual Bloom's Level
        q['blooms'] = identify_blooms_level(q.get('question', ''))
        
        # 3. Validate CO
        if q.get('co') not in available_co_ids:
            q['co'] = available_co_ids[(i + 2) % len(available_co_ids)]
            
        # 3.5 ENFORCE MARKS (User Requirement)
        q['marks'] = 10
        ans = q.get('answer', '')
        if isinstance(ans, dict):
            consolidated = ""
            for idx, (k, v) in enumerate(ans.items(), 1):
                # Clean prefix from both key and value if they match numeric patterns
                content = str(v).strip() if len(str(k)) < 30 else f"{k}: {v}"
                content = re.sub(rf"^{idx}\.?\s*{idx}\.?\s*", "", content, flags=re.IGNORECASE)
                content = re.sub(rf"^{idx}\.?\s*", "", content, flags=re.IGNORECASE)
                consolidated += f"{idx}. {content}\n"
            q['answer'] = consolidated.strip()
        elif isinstance(ans, list):
            consolidated = ""
            for idx, x in enumerate(ans, 1):
                content = str(x).strip()
                content = re.sub(rf"^{idx}\.?\s*", "", content, flags=re.IGNORECASE)
                consolidated += f"{idx}. {content}\n"
            q['answer'] = consolidated.strip()
        elif isinstance(ans, str):
            # Even if it's a string, ensure no triple numbering occurs
            lines = ans.split('\n')
            new_lines = []
            for idx, line in enumerate(lines, 1):
                clean_line = re.sub(rf"^{idx}\.?\s*{idx}\.?\s*", f"{idx}. ", line.strip(), flags=re.IGNORECASE)
                new_lines.append(clean_line)
            q['answer'] = "\n".join(new_lines)

    # Calculate Distributions
    co_marks = {}
    bl_marks = {l: 0 for l in BLOOMS_TAXONOMY.keys()}
    
    # Initialize with syllabus COs
    for c in cos:
        id_clean = c['id'].upper().replace(" ", "")
        co_marks[id_clean] = 0

    for q in paper.get('part_a', []):
        if not q or not isinstance(q, dict): continue
        co_raw = str(q.get('co', 'CO1')).strip().upper().replace(" ", "")
        # Handle cases like CO-1 or CO_1
        co = re.sub(r'[^A-Z0-9]', '', co_raw)
        bl = str(q.get('blooms', 'L1')).strip().upper()
        if co not in co_marks: co_marks[co] = 0
        co_marks[co] += 1
        if bl in bl_marks: bl_marks[bl] += 1
        
    for q in paper.get('part_b', []):
        if not q or not isinstance(q, dict): continue
        co_raw = str(q.get('co', 'CO1')).strip().upper().replace(" ", "")
        co = re.sub(r'[^A-Z0-9]', '', co_raw)
        bl = str(q.get('blooms', 'L1')).strip().upper()
        if co not in co_marks: co_marks[co] = 0
        co_marks[co] += 10 # Force 10 marks per question in Part B
        if bl in bl_marks: bl_marks[bl] += 10
        
    paper['distributions'] = {
        'co': co_marks,
        'blooms': bl_marks
    }
    paper['cos_definitions'] = cos
    
    return paper

# PDF and Word generation are now handled in formatting.py


def generate_chart_images(dist):
    """Generates CO and Blooms chart images and returns paths to temporary files."""
    co_img = tempfile.NamedTemporaryFile(suffix='.png', delete=False).name
    bl_img = tempfile.NamedTemporaryFile(suffix='.png', delete=False).name
    
    # 1. CO Bar Chart
    co_data = dist.get('co', {})
    if co_data:
        fig, ax = plt.subplots(figsize=(6, 4))
        ax.bar(co_data.keys(), co_data.values(), color='#1f77b4')
        ax.set_title('Course Outcome Vs Mark Distribution', fontweight='bold')
        ax.set_ylabel('Marks')
        for i, v in enumerate(co_data.values()):
            ax.text(i, v + 0.2, str(v), ha='center')
        plt.tight_layout()
        plt.savefig(co_img)
        plt.close()
    else:
        co_img = None
        
    # 2. Blooms Pie Chart
    bl_data = dist.get('blooms', {})
    if bl_data:
        fig, ax = plt.subplots(figsize=(7, 5)) # Slightly larger for pie
        sorted_keys = sorted(bl_data.keys())
        active_labels = [k for k in sorted_keys if bl_data[k] > 0]
        active_vals = [bl_data[k] for k in active_labels]
        if active_vals:
            # Use pctdistance and labeldistance to prevent overlapping
            # Also use a startangle for better distribution
            ax.pie(active_vals, 
                   labels=active_labels, 
                   autopct='%1.1f%%', 
                   colors=plt.cm.Paired.colors,
                   pctdistance=0.8,
                   labeldistance=1.1,
                   startangle=140)
            ax.set_title('Bloom\'s Level Vs Mark Distribution', fontweight='bold', pad=20)
            plt.tight_layout()
            plt.savefig(bl_img)
            plt.close()
        else:
            bl_img = None
    else:
        bl_img = None
        
    return co_img, bl_img


# Word generation is now handled in formatting.py
